"""Reads CGC_Dashboard_Content.xlsx and regenerates data.js.  Usage: python build_data.py"""
import json, sys
from openpyxl import load_workbook
XLSX = sys.argv[1] if len(sys.argv) > 1 else "CGC_Dashboard_Content.xlsx"
wb = load_workbook(XLSX, data_only=True)
def rows(name):
    return [r for r in wb[name].iter_rows(min_row=2, values_only=True) if any(c is not None for c in r)]
s = lambda v: "" if v is None else str(v)
n = lambda v: None if v in (None, "—", "") else v
funnel = rows("Funnel")
D = dict(
 screens=[dict(id=r[0], nav=r[1], title=r[2], desc=s(r[3]), lens=s(r[4]), filter=s(r[5])) for r in rows("Screens")],
 kpis=[dict(screen=r[0], title=r[1], value=s(r[2]), delta=s(r[3]), dir=s(r[4]), status=s(r[5]),
            tags=[t for t in s(r[6]).split("|") if t], note=s(r[7]), spark=int(r[8] or 0), progress=s(r[9])) for r in rows("KPIs")],
 funnel=[dict(label=r[0], value=r[1], sub=s(r[3]),
              pct="" if i == 0 else f"{funnel[i][1] / funnel[i-1][1] * 100:.1f}%") for i, r in enumerate(funnel)],
 map=[dict(code=r[0], value=n(r[1])) for r in rows("StateMap")],
 alerts=[dict(title=r[0], text=s(r[1]), owner=s(r[2]), due=s(r[3]), severity=s(r[4])) for r in rows("Alerts")],
 imfunnel=[dict(label=r[0], value=r[1], pct=s(r[2])) for r in rows("imSME_Funnel")],
 benchmarks=[dict(measure=r[0], current=s(r[1]), benchmark=s(r[2]), actual=s(r[3])) for r in rows("Benchmarks")],
 pfi=[dict(name=r[0], util=r[1], target=r[2], pct=r[3]) for r in rows("PFI")],
 gov=[dict(label=r[0], value=s(r[1]), note=s(r[2])) for r in rows("Governance")],
 exceptions=[dict(type=r[0], open=r[1], overdue=r[2]) for r in rows("Exceptions")],
 disclosure=[dict(pillar=r[0], pct=r[1], note=s(r[2])) for r in rows("Disclosure")],
 ops=[dict(label=r[0], value=s(r[1]), yoy=s(r[2])) for r in rows("OwnOps")],
 segments=[dict(label=r[0], pct=r[1], count=s(r[2])) for r in rows("Segments")],
)
open("data.js", "w", encoding="utf-8").write("window.DATA = " + json.dumps(D, ensure_ascii=False, indent=1) + ";\n")
print("data.js regenerated from", XLSX)
